"""Pricing workbook: Inputs (editable, blue), Per Car (measured slicer data + formulas), Price Guide, Sources.
usage: python make_pricing.py <out.xlsx> <inputs.json> <cars.json>"""
import json, sys
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.comments import Comment
from openpyxl.utils import get_column_letter

out, inputs_path, cars_path = sys.argv[1:4]
INP = json.load(open(inputs_path))          # {"groups": [{"title":..., "rows": [{"key","label","value","unit","fmt","note","source","key_assumption"}]}], "sources": [...]}
CARS = json.load(open(cars_path))           # list of per-car measured data

F = 'Arial'
BLUE = Font(name=F, color='0000FF'); BLACK = Font(name=F); BOLD = Font(name=F, bold=True)
GREEN = Font(name=F, color='008000')
H1 = Font(name=F, bold=True, size=14); H2 = Font(name=F, bold=True, size=11, color='FFFFFF')
HFILL = PatternFill('solid', fgColor='2F3B52'); YEL = PatternFill('solid', fgColor='FFFF00'); GREY = PatternFill('solid', fgColor='F2F2F2')
thin = Side(style='thin', color='BFBFBF'); BOX = Border(left=thin, right=thin, top=thin, bottom=thin)
USD = '$#,##0.00;($#,##0.00);-'; USD3 = '$#,##0.000;($#,##0.000);-'; PCT = '0.0%;(0.0%);-'; NUM1 = '#,##0.0;(#,##0.0);-'; NUM2 = '#,##0.00;(#,##0.00);-'
FMT = {'usd': USD, 'usd3': USD3, 'pct': PCT, 'num1': NUM1, 'num2': NUM2, 'int': '#,##0'}

wb = Workbook()
# ------------------------------------------------------------------ Inputs
ws = wb.active; ws.title = 'Inputs'
ws['A1'] = 'Keychain pricing model - inputs'; ws['A1'].font = H1
ws['A2'] = 'Blue cells are inputs: change them to your real costs. Yellow = key assumptions worth checking first. Black cells are formulas.'
ws['A2'].font = Font(name=F, italic=True)
ws.column_dimensions['A'].width = 46; ws.column_dimensions['B'].width = 14; ws.column_dimensions['C'].width = 14; ws.column_dimensions['D'].width = 90
REF = {}
r = 4
for g in INP['groups']:
    ws.cell(r, 1, g['title']).font = H2
    for c in range(1, 5):
        ws.cell(r, c).fill = HFILL
    ws.cell(r, 2, 'Value').font = H2; ws.cell(r, 3, 'Unit').font = H2; ws.cell(r, 4, 'Note / source').font = H2
    r += 1
    for row in g['rows']:
        ws.cell(r, 1, row['label']).font = BLACK
        v = row['value']
        cell = ws.cell(r, 2, v)
        is_formula = isinstance(v, str) and v.startswith('=')
        cell.font = BLACK if is_formula else BLUE
        cell.number_format = FMT.get(row.get('fmt', 'num2'), NUM2)
        if row.get('key_assumption'):
            cell.fill = YEL
        ws.cell(r, 3, row.get('unit', '')).font = BLACK
        note = row.get('note', '')
        if row.get('source'):
            note = (note + ' | ' if note else '') + 'Source: ' + row['source']
        ws.cell(r, 4, note).font = Font(name=F, size=9)
        ws.cell(r, 4).alignment = Alignment(wrap_text=True, vertical='top')
        for c in range(1, 5):
            ws.cell(r, c).border = BOX
        REF[row['key']] = f"Inputs!$B${r}"
        r += 1
    r += 1


def expand(expr):
    """Replace {key} with the absolute Inputs reference."""
    for k, a in REF.items():
        expr = expr.replace('{' + k + '}', a)
    return expr


# second pass: formulas inside Inputs that reference other inputs
for row_i in range(4, r):
    c = ws.cell(row_i, 2)
    if isinstance(c.value, str) and c.value.startswith('=') and '{' in c.value:
        c.value = expand(c.value)

# ------------------------------------------------------------------ Per Car
pc = wb.create_sheet('Per Car')
pc['A1'] = 'Cost per keychain, per car (slicer data measured on full K2 plates)'; pc['A1'].font = H1
pc['A2'] = ('Plate time / grams come from Creality Print 7.1 slices of the delivered plate files. 1-swap = recommended bulk build. '
            'Custom colour = 3-filament version (body colour of your choice), ~8 filament changes per plate instead of 1.')
pc['A2'].font = Font(name=F, italic=True)
cols = [
    ('Car', None, 30, None),
    ('Keychains per plate', 'n', 10, 'int'),
    ('1-swap plate time (h)', 'swap_h', 11, 'num2'),
    ('1-swap plate filament (g)', 'swap_g', 11, 'num1'),
    ('Classic plate time (h)', 'cl_h', 11, 'num2'),
    ('Classic plate filament (g)', 'cl_g', 11, 'num1'),
    ('Extra time per filament change (min)', 'chg_min', 12, 'num1'),
    ('Extra filament per change (g)', 'chg_g', 11, 'num2'),
]
hdr = 4
for j, (t, k, w, f) in enumerate(cols, 1):
    c = pc.cell(hdr, j, t); c.font = H2; c.fill = HFILL; c.alignment = Alignment(wrap_text=True, vertical='center')
    pc.column_dimensions[get_column_letter(j)].width = w
calc = [
    ('1-swap: filament $', '=D{r}/B{r}/1000*{filament_price_kg}*(1+{failure_rate})', 'usd3'),
    ('1-swap: machine $ (power + wear + depreciation)', '=C{r}/B{r}*{machine_cost_h}*(1+{failure_rate})', 'usd3'),
    ('Hardware $ (ring, chain, jump ring)', '={hardware_unit}', 'usd3'),
    ('Labour $', '=({labour_min_unit}+{labour_min_plate}/B{r})/60*{labour_rate}', 'usd3'),
    ('Packaging $', '={packaging_unit}', 'usd3'),
    ('1-swap: TOTAL cost per keychain', '=SUM(I{r}:M{r})', 'usd'),
    ('Classic: TOTAL cost per keychain', '=F{r}/B{r}/1000*{filament_price_kg}*(1+{failure_rate})+E{r}/B{r}*{machine_cost_h}*(1+{failure_rate})+K{r}+L{r}+M{r}', 'usd'),
    ('Custom colour (1-swap style): TOTAL cost', '=N{r}+{custom_extra_changes}*(H{r}/1000*{filament_price_kg}+G{r}/60*{machine_cost_h})/B{r}', 'usd'),
    ('Custom colour: plate time (h)', '=C{r}+{custom_extra_changes}*G{r}/60', 'num2'),
    ('Minutes of printer time per keychain (1-swap)', '=C{r}*60/B{r}', 'num1'),
    ('Plates per week (at {print_hours_week} h)', '={print_hours_week}/C{r}', 'num1'),
    ('Keychains per week', '=S{r}*B{r}', 'int'),
]
for j, (t, fml, f) in enumerate(calc, len(cols) + 1):
    c = pc.cell(hdr, j, t.replace('{print_hours_week}', '')); c.font = H2; c.fill = HFILL; c.alignment = Alignment(wrap_text=True, vertical='center')
    pc.column_dimensions[get_column_letter(j)].width = 13
pc.row_dimensions[hdr].height = 62
first = hdr + 1
for i, car in enumerate(CARS):
    rr = first + i
    vals = [car['car'], car['n'], car['swap_h'], car['swap_g'], car['cl_h'], car['cl_g'], car['chg_min'], car['chg_g']]
    for j, v in enumerate(vals, 1):
        c = pc.cell(rr, j, v); c.font = BLUE if j > 1 else BOLD; c.border = BOX
        if cols[j - 1][3]:
            c.number_format = FMT[cols[j - 1][3]]
    for j, (t, fml, f) in enumerate(calc, len(cols) + 1):
        c = pc.cell(rr, j, expand(fml.replace('{r}', str(rr)))); c.font = BLACK; c.number_format = FMT[f]; c.border = BOX
last = first + len(CARS) - 1
avg = last + 1
pc.cell(avg, 1, 'Average').font = BOLD
for j in range(2, len(cols) + len(calc) + 1):
    L = get_column_letter(j)
    c = pc.cell(avg, j, f'=AVERAGE({L}{first}:{L}{last})'); c.font = BOLD; c.fill = GREY
    c.number_format = pc.cell(first, j).number_format
pc.cell(avg + 2, 1, 'Measured data (blue) source: Creality Print 7.1 slicing engine, K2 0.4 nozzle, your 0.20 mm process, full plates of the delivered 3MF files. '
        'Classic = real two-colour slice; 1-swap = one-colour slice + one colour change at the measured per-change cost.').font = Font(name=F, size=9, italic=True)
AVG = {k: f"'Per Car'!${get_column_letter(j)}${avg}" for j, k in enumerate(
    ['car', 'n', 'swap_h', 'swap_g', 'cl_h', 'cl_g', 'chg_min', 'chg_g', 'fil', 'mach', 'hw', 'lab', 'pack', 'tot_swap', 'tot_classic', 'tot_custom', 'custom_h', 'min_unit', 'plates_wk', 'units_wk'], 1)}

# ------------------------------------------------------------------ Price Guide
pg = wb.create_sheet('Price Guide')
pg['A1'] = 'Price guide (per keychain, 1-swap build, average car)'; pg['A1'].font = H1
pg['A2'] = 'Prices are inputs (blue) on this sheet - try your own. Profit = price - all costs - platform fees - shipping you absorb.'
pg['A2'].font = Font(name=F, italic=True)
heads = ['Channel', 'Price per keychain', 'Units per order', 'Shipping charged to buyer (per order)', 'Cost per keychain (make)',
         'Platform / payment fees per keychain', 'Order costs per keychain (label, mailer, packing time)', 'Profit per keychain', 'Margin', 'Profit per printer-hour', 'Note']
widths = [34, 12, 10, 14, 12, 13, 14, 12, 9, 13, 60]
for j, (h, w) in enumerate(zip(heads, widths), 1):
    c = pg.cell(4, j, h); c.font = H2; c.fill = HFILL; c.alignment = Alignment(wrap_text=True, vertical='center')
    pg.column_dimensions[get_column_letter(j)].width = w
pg.row_dimensions[4].height = 58
for i, ch in enumerate(INP['channels']):
    rr = 5 + i
    pg.cell(rr, 1, ch['name']).font = BOLD
    for j, key in ((2, 'price'), (3, 'units'), (4, 'ship_charged')):
        c = pg.cell(rr, j, ch[key]); c.font = BLUE; c.number_format = USD if j != 3 else '#,##0'
    pg.cell(rr, 2).fill = YEL
    cost = AVG['tot_custom'] if ch.get('custom') else AVG['tot_swap']
    pg.cell(rr, 5, f'={cost}').font = GREEN
    fee = expand(ch['fee_formula']).replace('{price}', f'B{rr}').replace('{units}', f'C{rr}').replace('{ship}', f'D{rr}')
    pg.cell(rr, 6, '=' + fee)
    ship = expand(ch['ship_formula']).replace('{units}', f'C{rr}').replace('{ship}', f'D{rr}')
    pg.cell(rr, 7, '=' + ship)
    pg.cell(rr, 8, f'=B{rr}+IF(C{rr}>0,D{rr}/C{rr},0)-E{rr}-F{rr}-G{rr}')
    pg.cell(rr, 9, f'=IF(B{rr}>0,H{rr}/(B{rr}+IF(C{rr}>0,D{rr}/C{rr},0)),0)')
    pg.cell(rr, 10, f'=IF({AVG["min_unit"]}>0,H{rr}/({AVG["min_unit"]}/60),0)')
    pg.cell(rr, 11, ch.get('note', '')).alignment = Alignment(wrap_text=True, vertical='top')
    for j in range(1, 12):
        c = pg.cell(rr, j); c.border = BOX
        if j in (5, 6, 7, 8, 10):
            c.number_format = USD
        if j == 9:
            c.number_format = PCT
        if j >= 5 and j <= 10 and c.font != GREEN:
            c.font = BLACK if j != 5 else GREEN
        pg.cell(rr, 11).font = Font(name=F, size=9)
rr = 6 + len(INP['channels'])
pg.cell(rr, 1, 'Etsy price needed for a target profit (single keychain, free shipping)').font = BOLD
pg.cell(rr + 1, 1, 'Target profit per keychain').font = H2; pg.cell(rr + 1, 1).fill = HFILL
pg.cell(rr + 1, 2, 'Required price').font = H2; pg.cell(rr + 1, 2).fill = HFILL
for k, tgt in enumerate((4.0, 6.0, 8.0, 10.0)):
    q = rr + 2 + k
    c = pg.cell(q, 1, tgt); c.font = BLUE; c.number_format = USD; c.border = BOX
    # price = (make cost + order costs + target + listing + fixed processing) / (1 - % fees)
    f = expand(f"=({AVG['tot_swap']}+({{label}}+{{mailer}}+{{labour_min_order}}/60*{{labour_rate}})+A{q}+{{etsy_listing}}+{{etsy_proc_fixed}})"
               f"/(1-({{etsy_tx}}+{{etsy_proc_pct}}+{{etsy_offsite_share}}*{{etsy_offsite}}))")
    c = pg.cell(q, 2, f); c.number_format = USD; c.border = BOX; c.font = BLACK
rr = rr + 7
pg.cell(rr, 1, 'Recommendations').font = H1
for k, line in enumerate(INP['recommendations']):
    c = pg.cell(rr + 1 + k, 1, line); c.font = Font(name=F); c.alignment = Alignment(wrap_text=False)

# ------------------------------------------------------------------ Sources
so = wb.create_sheet('Sources')
so['A1'] = 'Sources and research notes'; so['A1'].font = H1
so.column_dimensions['A'].width = 40; so.column_dimensions['B'].width = 50; so.column_dimensions['C'].width = 90
for j, h in enumerate(['Item', 'Finding', 'Source'], 1):
    c = so.cell(3, j, h); c.font = H2; c.fill = HFILL
for i, s in enumerate(INP['sources']):
    so.cell(4 + i, 1, s['item']).font = BLACK
    so.cell(4 + i, 2, s['finding']).font = BLACK
    so.cell(4 + i, 3, s['url']).font = Font(name=F, size=9, color='0563C1')
    for j in range(1, 4):
        so.cell(4 + i, j).alignment = Alignment(wrap_text=True, vertical='top')
for sh in wb.worksheets:
    for row in sh.iter_rows():
        for c in row:
            if c.font and c.font.name != F:
                c.font = Font(name=F, bold=c.font.bold, italic=c.font.italic, size=c.font.size, color=c.font.color)
    sh.freeze_panes = None
pc.freeze_panes = 'B5'
wb.calculation.fullCalcOnLoad = True
wb.save(out)
print('saved', out)
